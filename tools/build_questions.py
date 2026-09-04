import json
import os
import re
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "KATALOG_072026.xlsx")
OUT = os.path.join(ROOT, "data", "questions.json")
OUT_JS = os.path.join(ROOT, "data", "questions.js")

CATEGORY = "B"

COL = {
    "num": 1,
    "q_pl": 2,
    "a_pl": 3,
    "b_pl": 4,
    "c_pl": 5,
    "correct": 6,
    "media": 7,
    "scope": 8,
    "points": 9,
    "cats": 10,
    "q_en": 15,
    "a_en": 16,
    "b_en": 17,
    "c_en": 18,
    "q_de": 19,
    "a_de": 20,
    "b_de": 21,
    "c_de": 22,
    "q_ua": 23,
    "a_ua": 24,
    "b_ua": 25,
    "c_ua": 26,
}

LANGS = ["pl", "en", "de", "ua"]


def clean(v):
    if v is None:
        return ""
    s = str(v)
    s = s.replace(" ", " ").replace("_x000D_", "")
    s = unicodedata.normalize("NFC", s)
    return re.sub(r"\s+", " ", s).strip()


def scope_of(raw):
    s = clean(raw).lower()
    if s.startswith("pod"):
        return "P"
    if s.startswith("spec") or s.startswith("specaj"):
        return "S"
    return None


def media_of(raw):
    name = clean(raw)
    if not name:
        return None
    base, ext = os.path.splitext(name)
    ext = ext.lower()
    if ext == ".wmv":
        return {"kind": "video", "src": base + ".mp4", "orig": name}
    if ext in (".jpg", ".jpeg", ".png"):
        return {"kind": "image", "src": name, "orig": name}
    return {"kind": "image", "src": name, "orig": name}


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    out = []
    seen = set()
    skipped = {"no_num": 0, "not_cat": 0, "no_scope": 0, "bad_answer": 0, "dupe": 0}

    for sheet in wb.sheetnames:
        verified = sheet.strip().lower() == "katalog"
        rows = wb[sheet].iter_rows(values_only=True)
        next(rows)
        for r in rows:
            num = clean(r[COL["num"]])
            if not num:
                skipped["no_num"] += 1
                continue
            cats = [c.strip() for c in clean(r[COL["cats"]]).split(",") if c.strip()]
            if CATEGORY not in cats:
                skipped["not_cat"] += 1
                continue
            scope = scope_of(r[COL["scope"]])
            if scope is None:
                skipped["no_scope"] += 1
                continue
            if num in seen:
                skipped["dupe"] += 1
                continue

            correct = clean(r[COL["correct"]]).upper()
            opts_pl = [clean(r[COL[k + "_pl"]]) for k in ("a", "b", "c")]
            has_opts = any(opts_pl)

            if has_opts:
                if correct not in ("A", "B", "C"):
                    skipped["bad_answer"] += 1
                    continue
                kind = "abc"
            else:
                if correct not in ("T", "N"):
                    skipped["bad_answer"] += 1
                    continue
                kind = "tn"

            try:
                points = int(clean(r[COL["points"]]))
            except ValueError:
                skipped["bad_answer"] += 1
                continue

            q = {}
            opts = {}
            for lang in LANGS:
                q[lang] = clean(r[COL["q_" + lang]])
                if kind == "abc":
                    trio = [clean(r[COL[k + "_" + lang]]) for k in ("a", "b", "c")]
                    if any(trio):
                        opts[lang] = trio

            for lang in LANGS:
                if not q[lang]:
                    q[lang] = q["pl"]
                if kind == "abc" and lang not in opts:
                    opts[lang] = opts.get("pl", ["", "", ""])

            seen.add(num)
            item = {
                "id": num,
                "scope": scope,
                "points": points,
                "kind": kind,
                "correct": correct,
                "q": q,
                "verified": verified,
            }
            if kind == "abc":
                item["opts"] = opts
            media = media_of(r[COL["media"]])
            if media:
                item["media"] = media
            out.append(item)

    out.sort(key=lambda x: int(x["id"]) if x["id"].isdigit() else 0)

    stats = {
        "total": len(out),
        "podstawowy": sum(1 for q in out if q["scope"] == "P"),
        "specjalistyczny": sum(1 for q in out if q["scope"] == "S"),
        "video": sum(1 for q in out if q.get("media", {}).get("kind") == "video"),
        "image": sum(1 for q in out if q.get("media", {}).get("kind") == "image"),
        "no_media": sum(1 for q in out if "media" not in q),
        "unverified": sum(1 for q in out if not q["verified"]),
        "skipped": skipped,
    }
    for scope in ("P", "S"):
        for pts in (1, 2, 3):
            key = "pool_%s%d" % (scope, pts)
            stats[key] = sum(1 for q in out if q["scope"] == scope and q["points"] == pts)

    payload = {"category": CATEGORY, "langs": LANGS, "questions": out}
    blob = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    with open(OUT, "w", encoding="utf-8") as f:
        f.write(blob)

    # file:// pages cannot fetch JSON, so ship the same payload as a script.
    with open(OUT_JS, "w", encoding="utf-8") as f:
        f.write("window.KATALOG=")
        f.write(blob)
        f.write(";\n")

    json.dump(stats, sys.stdout, ensure_ascii=False, indent=2)
    print()
    print("wrote %s (%.1f MB)" % (OUT, os.path.getsize(OUT) / 1048576.0))


if __name__ == "__main__":
    main()
